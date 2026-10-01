# -*- coding: utf-8 -*-
# ==============================<seiki>‌​‌​‌​‍‌‌​​‎‌​‍‎‌‌‎‎‌​​‎‌​‎‎‌​‌​‌​‌‌‌‌‍​‌‌‎‎‌‌​‎‌​‍‌‌​‌‎‌‌‎‎​‎‌‎‌‍‌‍​‎​‎‌‍​‌​‎‍‌‌‍​‎​‎​‍‌‍‌‌==============================
#   novel_export.py —— 把树状聊天记录线性化成小说（Word / EPUB）
#
#   和现有「导出 Word」的区别：
#     现有那份导的是 self.messages（扁平消息列表），分支、重生成、回溯过的内容
#     全糊在一起，读起来不像故事。
#     这里走【树上你实际走过的路径】（root_id → current_leaf_id），
#     并把当时没选的那条分支收成脚注 —— 于是每局对话都能变成一部能读的作品。
#
#   EPUB 是手写的（zipfile + XML），不引第三方库：
#     EPUB 本质就是一个 zip，规定 mimetype 必须是【第一个条目且不压缩】。
#     自己写反而更好控制（项目风格也是能零依赖就零依赖）。
# ============================================================

import io
import os
import re
import uuid
import zipfile
from datetime import datetime

EPUB_MIME = "application/epub+zip"
_CHAPTER_MAX_BLOCKS = 25


# ---------- 小工具 ----------
def xml_escape(s):
    """XML 文本转义。聊天里出现 & < > 是常事，不转义直接就是坏 XHTML。"""
    return (str(s if s is not None else "")
            .replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


def _clip(s, n):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[:n - 1] + "…"


def _is_sys(node):
    return str(node.get("role") or "").lower() in ("sys", "system")


def _speaker_of(node):
    meta = node.get("metadata") or {}
    who = str(meta.get("speaker") or "").strip()
    if who:
        return who
    if str(node.get("role") or "").lower() == "user":
        return "你"
    return "AI"


def _ts_key(node):
    return str(node.get("timestamp") or "")


# ---------- 线性化 ----------
def resolve_leaf(tree):
    """决定"当前路径"的终点。

    优先用树自己记的 current_leaf_id；它缺失/失效时，退回到"时间戳最新的叶子"，
    全空则用 root。绝不返回不存在的 id —— 否则线性化会整段丢掉。
    """
    nodes = (tree or {}).get("nodes") or {}
    leaf = (tree or {}).get("current_leaf_id")
    if leaf and leaf in nodes:
        return leaf
    if not nodes:
        return None
    leaves = [nid for nid, n in nodes.items() if not (n.get("children_ids") or [])]
    pool = leaves or list(nodes.keys())
    return max(pool, key=lambda nid: _ts_key(nodes[nid]))


def path_to_root(tree, leaf):
    """从叶子沿 parent_id 回溯到根，返回正序列表。带环保护。"""
    nodes = (tree or {}).get("nodes") or {}
    out = []
    seen = set()
    cur = leaf
    while cur and cur in nodes and cur not in seen:
        seen.add(cur)
        out.append(nodes[cur])
        cur = nodes[cur].get("parent_id")
    out.reverse()
    return out


def _chapter_excerpt(content, n=16):
    """从一句发言里取干净的章节小标题。

    真实存档里用户发言常带说话人标记（`[Echo] 怎么？你喜欢我？`）或全角舞台提示
    （`（心里生成了一个鬼点子）…`），直接拿来做标题很脏，这里先洗掉再截断；
    洗不出东西就返回空串 —— 那样标题只留编号，也比挂个 `[Echo]` 像书。
    """
    s = re.sub(r"\s+", " ", str(content or "")).strip()
    for _ in range(2):
        s = re.sub(r"^\s*[\[【][^\]】]{1,20}[\]】]\s*", "", s)
        s = re.sub(r"^\s*[（(][^）)]{0,40}[）)]\s*", "", s)
    s = s.strip(" ：:、,，。.…~-—")
    if not s:
        return ""
    return s if len(s) <= n else s[:n - 1] + "…"


def linearize(tree, include_sys=False, max_alt=140, chapter_size=3):
    """把树线性化成小说结构。

    返回 {"ok", "meta", "chapters": [...], "stats": {...}}
    chapters[i] = {"title", "blocks": [{"type","speaker","text"}], "alts": [str, ...]}
    block.type ∈ dialogue / note / image / sys

    chapter_size：每几个「用户回合」切一章。默认 3 —— 按每条用户发言切会让对话体
    小说碎成一段一段（实测真实存档 14 个块切出 7 章，读起来像流水账）。
    """
    nodes = (tree or {}).get("nodes") or {}
    leaf = resolve_leaf(tree)
    if not nodes or not leaf:
        return {"ok": False, "msg": "这局还没有对话内容，导不出小说", "meta": {}, "chapters": [],
                "stats": {}}
    nodes_path = path_to_root(tree, leaf)
    if not include_sys:
        nodes_path = [n for n in nodes_path if not _is_sys(n)]
    if not nodes_path:
        return {"ok": False, "msg": "路径上只剩系统消息，导不出小说", "meta": {}, "chapters": [],
                "stats": {}}

    on_path = set(id(n) for n in nodes_path)
    chapters = []
    stats = {"chapters": 0, "blocks": 0, "words": 0, "branch_points": 0,
             "alts": 0, "path_nodes": len(nodes_path),
             "off_path_nodes": sum(1 for n in nodes.values() if id(n) not in on_path)}
    cur = None
    users_in_chapter = 0

    def new_chapter(excerpt="", prologue=False):
        if prologue:
            title = u"序"
        else:
            num = sum(1 for c in chapters if not c["title"].startswith(u"序")) + 1
            title = u"第 %d 章" % num
            if excerpt:
                title += u" · " + excerpt
        ch = {"title": title, "blocks": [], "alts": []}
        chapters.append(ch)
        return ch

    for idx, node in enumerate(nodes_path):
        role = str(node.get("role") or "").lower()
        content = str(node.get("content") or "").strip()
        meta = node.get("metadata") or {}
        next_id = nodes_path[idx + 1].get("id") if idx + 1 < len(nodes_path) else None
        block = None
        if _is_sys(node):
            block = {"type": "sys", "speaker": "系统", "text": content}
        elif meta.get("image") or meta.get("image_url"):
            block = {"type": "image", "speaker": _speaker_of(node),
                     "text": content or "（图片）"}
        elif role == "user":
            block = {"type": "dialogue", "speaker": "你", "text": content}
        else:
            block = {"type": "dialogue", "speaker": _speaker_of(node), "text": content}

        # 切章：开场白单独进「序」；从第一条用户发言起编号切章（每 chapter_size 个回合）；
        # 单章过长也切。按每条发言切会让对话体小说碎掉（实测 14 块切 7 章）。
        if role == "user":
            if cur is None or cur["title"] == u"序":
                cur = new_chapter(_chapter_excerpt(content))
                users_in_chapter = 0
            elif users_in_chapter >= chapter_size:
                cur = new_chapter(_chapter_excerpt(content))
                users_in_chapter = 0
            users_in_chapter += 1
        elif cur is None:
            cur = new_chapter(prologue=True)
        elif len(cur["blocks"]) >= _CHAPTER_MAX_BLOCKS:
            cur = new_chapter(_chapter_excerpt(content))
            users_in_chapter = 0
        cur["blocks"].append(block)
        stats["blocks"] += 1
        stats["words"] += len(content)

        # 分支点：当时还选了别的 → 收成脚注（只取没走的那几个孩子本身）
        kids = node.get("children_ids") or []
        if len(kids) > 1:
            stats["branch_points"] += 1
            for kid in kids:
                if kid == next_id:
                    continue
                alt = nodes.get(kid)
                if not alt:
                    continue
                alt_text = _clip(alt.get("content"), max_alt)
                if alt_text:
                    cur["alts"].append(u"另一条路：%s" % alt_text)
                    stats["alts"] += 1

    stats["chapters"] = len(chapters)
    secs = stats["words"] / 5.0
    stats["est_minutes"] = round(secs / 60.0, 1)
    meta = {"title": (tree or {}).get("name") or "", "leaf": leaf,
            "generated": datetime.now().strftime("%Y-%m-%d %H:%M")}
    return {"ok": True, "msg": "ok", "meta": meta, "chapters": chapters, "stats": stats}


# ---------- Word ----------
def to_docx(novel, path, title="", author=""):
    """写出 .docx（python-docx 是项目必需依赖，不算新增）。"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(title or novel["meta"].get("title") or u"聊天记录·小说版", level=0)
    if author:
        p = doc.add_paragraph()
        r = p.add_run(u"作者：%s" % author)
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    st = novel.get("stats") or {}
    p = doc.add_paragraph()
    r = p.add_run(u"%s 生成 · %d 章 · 约 %s 字 · 预估阅读 %s 分钟"
                  % (novel["meta"].get("generated", ""), st.get("chapters", 0),
                     st.get("words", 0), st.get("est_minutes", 0)))
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    for ch in novel["chapters"]:
        doc.add_heading(ch["title"], level=1)
        for b in ch["blocks"]:
            if b["type"] == "sys":
                p = doc.add_paragraph()
                r = p.add_run(u"〔%s〕" % b["text"])
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            elif b["type"] == "image":
                p = doc.add_paragraph()
                r = p.add_run(u"〔图片〕%s" % b["text"])
                r.italic = True
            else:
                p = doc.add_paragraph()
                rr = p.add_run(b["speaker"] + "：")
                rr.bold = True
                p.add_run(b["text"])
        for alt in ch["alts"]:
            p = doc.add_paragraph()
            r = p.add_run(u"※ " + alt)
            r.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    doc.save(path)
    return True


# ---------- EPUB（手写，零依赖） ----------
_CONTAINER = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""

_CSS = """body { font-family: "Noto Serif CJK SC", serif; line-height: 1.8; margin: 0 6%; }
h1 { font-size: 1.3em; margin: 1.6em 0 0.8em; }
p { margin: 0.5em 0; text-indent: 0; }
p.sys, p.alt, p.meta { color: #888; font-size: 0.86em; font-style: italic; }
p.meta { border-bottom: 1px solid #ddd; padding-bottom: 0.6em; }
span.sp { font-weight: bold; }
"""


def _xhtml(title, body):
    return ('<?xml version="1.0" encoding="UTF-8"?>\n'
            '<!DOCTYPE html>\n'
            '<html xmlns="http://www.w3.org/1999/xhtml" '
            'xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="zh" lang="zh">\n'
            '<head><meta charset="utf-8"/><title>%s</title>'
            '<link rel="stylesheet" type="text/css" href="styles.css"/></head>\n'
            '<body>\n%s\n</body>\n</html>\n' % (xml_escape(title), body))


def to_epub(novel, path, title="", author="", lang="zh"):
    """写出 EPUB 3（mimetype 必须第一个且不压缩，这是规范硬要求）。"""
    book_title = title or novel["meta"].get("title") or u"聊天记录·小说版"
    st = novel.get("stats") or {}
    uid = "urn:uuid:" + str(uuid.uuid4())
    chapters = novel["chapters"] or [{"title": u"空", "blocks": [], "alts": []}]

    pages = []
    for i, ch in enumerate(chapters, 1):
        parts = ["<h1>%s</h1>" % xml_escape(ch["title"])]
        for b in ch["blocks"]:
            if b["type"] == "sys":
                parts.append('<p class="sys">〔%s〕</p>' % xml_escape(b["text"]))
            elif b["type"] == "image":
                parts.append('<p class="sys">〔图片〕%s</p>' % xml_escape(b["text"]))
            else:
                parts.append('<p><span class="sp">%s：</span>%s</p>'
                             % (xml_escape(b["speaker"]), xml_escape(b["text"])))
        for alt in ch["alts"]:
            parts.append('<p class="alt">※ %s</p>' % xml_escape(alt))
        pages.append(("chapter%d.xhtml" % i, u"%s" % ch["title"], "\n".join(parts)))

    nav_items = "\n".join(
        '<li><a href="%s">%s</a></li>' % (fn, xml_escape(t)) for fn, t, _b in pages)
    nav = _xhtml(u"目录", '<nav epub:type="toc" id="toc"><h1>目录</h1><ol>%s</ol></nav>'
                 % nav_items)
    ncx_items = "\n".join(
        '<navPoint id="n%d" playOrder="%d"><navLabel><text>%s</text></navLabel>'
        '<content src="%s"/></navPoint>' % (i, i, xml_escape(t), fn)
        for i, (fn, t, _b) in enumerate(pages, 1))
    ncx = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
           '<head><meta name="dtb:uid" content="%s"/></head>'
           '<docTitle><text>%s</text></docTitle><navMap>%s</navMap></ncx>'
           % (xml_escape(uid), xml_escape(book_title), ncx_items))

    manifest = "\n".join(
        '<item id="c%d" href="%s" media-type="application/xhtml+xml"/>' % (i, fn)
        for i, (fn, _t, _b) in enumerate(pages, 1))
    spine = "\n".join('<itemref idref="c%d"/>' % i for i in range(1, len(pages) + 1))
    opf = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bid">\n'
           '  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
           '    <dc:identifier id="bid">%s</dc:identifier>\n'
           '    <dc:title>%s</dc:title>\n'
           '    <dc:creator>%s</dc:creator>\n'
           '    <dc:language>%s</dc:language>\n'
           '    <meta property="dcterms:modified">%s</meta>\n'
           '    <meta name="generator" content="Direct-Interface Cork-bore Kit"/>\n'
           '  </metadata>\n'
           '  <manifest>\n'
           '    <item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
           '    <item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
           '    <item id="css" href="styles.css" media-type="text/css"/>\n'
           '    %s\n'
           '  </manifest>\n'
           '  <spine toc="ncx">\n    %s\n  </spine>\n'
           '</package>\n'
           % (xml_escape(uid), xml_escape(book_title), xml_escape(author or ""),
              xml_escape(lang), datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"), manifest, spine))

    tmp = path + ".part"
    try:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            # 规范：mimetype 必须第一个条目，且不得压缩
            zi = zipfile.ZipInfo("mimetype", date_time=datetime.now().timetuple()[:6])
            zi.compress_type = zipfile.ZIP_STORED
            z.writestr(zi, EPUB_MIME)
            z.writestr("META-INF/container.xml", _CONTAINER)
            z.writestr("OEBPS/content.opf", opf)
            z.writestr("OEBPS/nav.xhtml", nav)
            z.writestr("OEBPS/toc.ncx", ncx)
            z.writestr("OEBPS/styles.css", _CSS)
            for fn, t, body in pages:
                z.writestr("OEBPS/" + fn, _xhtml(t, body))
        os.replace(tmp, path)
    except Exception:
        # 失败绝不能留下 .part 垃圾：目标被占用（例如阅读器开着这本书）时
        # os.replace 会抛 PermissionError，这里要收拾干净再抛。
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass
        raise
    return True


def safe_name(s, fallback="novel"):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", str(s or "")).strip(" .")
    return s[:60] or fallback


def export_novel(tree, fmt, path, title="", author="", include_sys=False):
    """一站式：线性化 + 写文件。返回 (ok, 消息)"""
    nov = linearize(tree, include_sys=include_sys)
    if not nov.get("ok"):
        return False, nov.get("msg") or "没有可导出的内容"
    try:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        if str(fmt).lower() in ("epub",):
            to_epub(nov, path, title=title, author=author)
        elif str(fmt).lower() in ("docx", "word"):
            to_docx(nov, path, title=title, author=author)
        else:
            return False, "不支持的格式：%s（只支持 docx / epub）" % fmt
    except Exception as e:
        return False, "导出失败：%s" % str(e)[:160]
    st = nov["stats"]
    return True, ("已导出 %d 章 · 约 %s 字 · 预估 %s 分钟（脚注 %d 条分支）"
                  % (st["chapters"], st["words"], st["est_minutes"], st["alts"]))
